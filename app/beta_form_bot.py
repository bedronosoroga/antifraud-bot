from __future__ import annotations

import asyncio
import logging
import os
import sqlite3
import json
import csv
import tempfile
from sqlalchemy.exc import IntegrityError
from pathlib import Path
from typing import Any, Optional

from aiogram import Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.filters import CommandStart, Command
from aiogram.fsm.context import FSMContext
from aiogram.types import (
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    Message,
    ChatJoinRequest,
    ForceReply,
    FSInputFile,
)
from datetime import datetime, timezone
from dotenv import load_dotenv
from aiogram.exceptions import TelegramBadRequest
from app.beta_form_filters import ConfirmFilter
from app.beta_form_q4 import Q4_OPTIONS, render_q4_text, build_q4_kb

from app.beta_form_texts import (
    greeting,
    thanks,
    accept_text,
    bot_link_text,
    decline_text,
    reminder_start_text,
    reminder_3h_text,
    reminder_24h_text,
    reminder_48h_text,
)
from app.core import db as dal
from app.beta.whitelist import add_to_whitelist, is_whitelisted
from app.config import cfg
from app.domain.quotas.service import QuotaService
from app.beta.activity import (
    init_activity_db,
    ensure_activity,
    mark_start_reminder_sent,
    mark_3h_sent,
    mark_24h_sent,
    mark_48h_sent,
    mark_inactive,
    get_activity_rows,
    get_history_stats,
    now_utc,
)
from sqlalchemy import select, func
try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None
from datetime import time, timedelta
from functools import lru_cache

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(name)s: %(message)s")
logger = logging.getLogger("beta_form_bot")

load_dotenv()


def env_set_int(name: str) -> set[int]:
    raw = os.getenv(name, "") or ""
    items = [part.strip() for part in raw.replace(";", ",").split(",") if part.strip()]
    result: set[int] = set()
    for it in items:
        try:
            result.add(int(it))
        except ValueError:
            continue
    return result


BOT_TOKEN = os.getenv("BETA_FORM_BOT_TOKEN")
if not BOT_TOKEN:
    raise RuntimeError("BETA_FORM_BOT_TOKEN is required for the beta form bot")

TARGET_CHAT = os.getenv("BETA_FORM_CHAT_ID")
ADMIN_IDS = env_set_int("ADMIN_IDS")
GROUP_LINK = os.getenv("BETA_GROUP_LINK") or ""
TEST_BOT_LINK = os.getenv("BETA_TEST_BOT_LINK") or ""
GROUP_ID_RAW = os.getenv("BETA_GROUP_ID")
GROUP_ID = int(GROUP_ID_RAW) if GROUP_ID_RAW and GROUP_ID_RAW.lstrip("-").isdigit() else None
DB_PATH = Path(os.getenv("BETA_FORM_DB", "data/beta_forms.db"))
GRANT_QTY = int(os.getenv("BETA_GRANT_QTY", "30") or 30)

quota_service: QuotaService | None = None
db_conn: sqlite3.Connection | None = None

router = Router(name="beta_form")

STEP_ORDER = ["q1", "q2", "q3", "q4", "q5", "q5_extra", "q6", "q7", "q8", "q9", "confirm"]

BTN = InlineKeyboardButton


def _kb(rows):
    return InlineKeyboardMarkup(inline_keyboard=rows)


async def _notify_admins(bot: Bot, text: str) -> None:
    for admin_id in ADMIN_IDS:
        try:
            await bot.send_message(admin_id, text)
        except Exception:
            continue


# --- Survey definitions ----------------------------------------------------

# Интервалы рассылки опросов
SURVEY_MID_DELAY_HOURS = 72
SURVEY_FINAL_DELAY_HOURS = 168

SURVEY_QUESTIONS = {
    "mid": [
        {
            "code": "m1",
            "text": "1️⃣ <b>Насколько вам было легко разобраться с ботом в первый раз?</b>\n\nОцените по шкале, где 1 — совсем непонятно, 5 — всё просто.",
            "type": "scale5",
            "labels": ["1", "2", "3", "4", "5"],
            "required": True,
        },
        {
            "code": "m2",
            "text": "2️⃣ <b>Насколько понятен сам результат проверки?</b>\n\nПонимаете ли вы, что означают индексы и вывод по рискам?\nОцените по шкале, где 1 — ничего не понял(а), 5 — всё понятно.",
            "type": "scale5",
            "labels": ["1", "2", "3", "4", "5"],
            "required": True,
        },
        {
            "code": "m3",
            "text": "3️⃣ <b>Как бот работает по скорости?</b>",
            "type": "options",
            "options": ["Очень быстро", "Нормально", "Медленно, иногда раздражает"],
            "required": True,
        },
        {
            "code": "m4",
            "text": "4️⃣ <b>Вы уже успели применить бот в реальной задаче?</b>\n\nНе тест, а живой контрагент / конкретная сделка.",
            "type": "options",
            "options": ["Да, использую в работе", "Пока только пробовал(а)", "Ещё не успел(а)"],
            "required": True,
        },
        {
            "code": "m4_extra",
            "text": "4️⃣ <b>Если хотите, напишите 1–2 фразы:</b>\n\nВ каких ситуациях бот оказался полезен / не помог.",
            "type": "text_optional",
            "required": False,
        },
        {
            "code": "m5",
            "text": "5️⃣ <b>Что в боте было непонятнее всего за эти дни?</b>\n\nНапример: формулировка, кнопка, раздел, результат.\nЕсли нечего вспомнить — так и напишите: «без сложностей».",
            "type": "text",
            "required": True,
        },
        {
            "code": "m6",
            "text": "6️⃣ <b>Если бы можно было исправить только одну вещь — что бы выбрали?</b>",
            "type": "text_optional",
            "required": False,
        },
    ],
    "final": [
        {
            "code": "f1",
            "text": "1️⃣ <b>Насколько в целом вам зашёл бот?</b>\n\nОцените по шкале, где 1 — совсем не зашёл, 5 — очень нравится.",
            "type": "scale5",
            "labels": ["1", "2", "3", "4", "5"],
            "required": True,
        },
        {
            "code": "f2",
            "text": "2️⃣ <b>Насколько бот полезен именно в вашей работе?</b>",
            "type": "options",
            "options": ["Очень помогает, буду пользоваться", "Иногда полезен, но не критично", "Пока не чувствую реальной пользы"],
            "required": True,
        },
        {
            "code": "f2_extra",
            "text": "2️⃣ <b>Если есть пример, когда бот помог или не помог — напишите пару слов.</b>",
            "type": "text_optional",
            "required": False,
        },
        {
            "code": "f3",
            "text": "3️⃣ <b>Насколько понятно, что означают индексы и уровни риска?</b>\n\nОцените по шкале, где 1 — совсем не понятно, 5 — всё ясно.",
            "type": "scale5",
            "labels": ["1", "2", "3", "4", "5"],
            "required": True,
        },
        {
            "code": "f3_extra",
            "text": "3️⃣ <b>Что именно сбивает с толку?</b>\n\nНапример: «не понимаю, что считать хорошим индексом», «незапомнились уровни риска» и т. д.",
            "type": "text_optional",
            "required": False,
        },
        {
            "code": "f4",
            "text": "4️⃣ <b>Читали ли вы раздел про методику расчёта индексов?</b>",
            "type": "options",
            "options": ["Да, читал(а) внимательно", "Пробежался(ась) по диагонали", "Нет, не открывал(а)"],
            "required": True,
        },
        {
            "code": "f4_rate",
            "text": "4️⃣ <b>Насколько понятным оказался текст про методику и откуда берутся индексы?</b>\n\nОцените по пятибальной шкале, где 1 — ничего не понял(а), 5 — всё понятно.",
            "type": "scale5",
            "labels": ["1", "2", "3", "4", "5"],
            "required": False,
        },
        {
            "code": "f4_extra",
            "text": "4️⃣ <b>Что бы вы упростили или переформулировали в этом разделе?</b>\n\nМожно указать фразу или блок.",
            "type": "text_optional",
            "required": False,
        },
        {
            "code": "f5",
            "text": "5️⃣ <b>Насколько удобно пользоваться ботом в целом?</b>",
            "type": "options",
            "options": ["Очень удобно, всё логично", "В целом нормально, но есть шероховатости", "Скорее неудобно, часто путаюсь"],
            "required": True,
        },
        {
            "code": "f5_extra",
            "text": "5️⃣ <b>Если есть, напишите, какие шаги или кнопки кажутся лишними или запутанными.</b>",
            "type": "text_optional",
            "required": False,
        },
        {
            "code": "f6",
            "text": "6️⃣ <b>Что в боте вам больше всего понравилось?</b>\n\n1–2 пункта.",
            "type": "text_optional",
            "required": False,
        },
        {
            "code": "f6b",
            "text": "6️⃣ <b>Что показалось лишним или раздражающим?</b>",
            "type": "text_optional",
            "required": False,
        },
        {
            "code": "f7",
            "text": "7️⃣ <b>Если хотите — напишите всё, что ещё важно сказать.</b>\n\nЛюбые мысли, идеи, предложения.",
            "type": "text_optional",
            "required": False,
        },
    ],
}


@lru_cache(maxsize=None)
def _survey_steps(kind: str) -> list[dict[str, Any]]:
    return SURVEY_QUESTIONS.get(kind, [])


def _survey_step(kind: str, code: str) -> Optional[dict[str, Any]]:
    for step in _survey_steps(kind):
        if step["code"] == code:
            return step
    return None


def _survey_first(kind: str) -> Optional[str]:
    steps = _survey_steps(kind)
    return steps[0]["code"] if steps else None


def _survey_next(kind: str, current: Optional[str], answers: dict[str, Any]) -> Optional[str]:
    steps = _survey_steps(kind)
    if current is None:
        return steps[0]["code"] if steps else None
    codes = [s["code"] for s in steps]
    # conditional skips
    if kind == "mid" and current == "m4":
        val = answers.get("m4", "")
        if val == "Ещё не успел(а)":
            # пропускаем m4_extra
            try:
                idx = codes.index("m4_extra")
                next_code = codes[idx + 1] if idx + 1 < len(codes) else None
                return next_code
            except ValueError:
                pass
        else:
            return "m4_extra"
    if kind == "final" and current == "f3":
        val = answers.get("f3", "")
        if val in {"4", "5"}:
            # пропускаем уточняющий текст
            try:
                idx = codes.index("f3_extra")
                next_code = codes[idx + 1] if idx + 1 < len(codes) else None
                return next_code
            except ValueError:
                pass
    if kind == "final" and current == "f4":
        val = answers.get("f4", "")
        if val == "Нет, не открывал(а)":
            return "f4_extra"  # спрашиваем только открытый вопрос, без оценки текста
    try:
        idx = codes.index(current)
    except ValueError:
        return None
    if idx + 1 < len(steps):
        return steps[idx + 1]["code"]
    return None


def _survey_kb(kind: str, code: str, step: dict[str, Any]) -> Optional[InlineKeyboardMarkup]:
    buttons: list[list[InlineKeyboardButton]] = []
    if step["type"] == "scale5":
        btns = []
        labels = step.get("labels") or ["1", "2", "3", "4", "5"]
        for idx, lab in enumerate(labels):
            btns.append(InlineKeyboardButton(text=lab, callback_data=f"sv:{kind}:{code}:{idx}"))
        buttons.append(btns)
    elif step["type"] == "options":
        opts = step.get("options", [])
        for idx, opt in enumerate(opts):
            buttons.append([InlineKeyboardButton(text=opt, callback_data=f"sv:{kind}:{code}:{idx}")])
    elif step["type"] == "text_optional":
        buttons.append([InlineKeyboardButton(text="Пропустить", callback_data=f"svskip:{kind}:{code}")])
    if not buttons:
        return None
    return InlineKeyboardMarkup(inline_keyboard=buttons)


async def _send_survey_step(bot: Bot, uid: int, kind: str, code: str, answers: dict[str, Any]) -> int:
    step = _survey_step(kind, code)
    if not step:
        return 0
    text = step["text"]
    kb = _survey_kb(kind, code, step)
    force_new_msg = True  # в опросах всегда создаём новое, старое удаляем при ответе
    markup = None
    if step["type"] in {"text", "text_optional"}:
        markup = ForceReply(selective=False)
    msg = await bot.send_message(uid, text, reply_markup=kb or markup)
    _set_survey_state(uid, kind, current_step=code, message_id=msg.message_id)
    return msg.message_id


async def _start_survey(bot: Bot, uid: int, kind: str) -> bool:
    existing = _get_survey(uid, kind)
    if existing and existing["status"] == "answered":
        return False
    first = _survey_first(kind)
    if not first:
        return False
    _create_survey(uid, kind, status="sent", step=first, message_id=None)
    if kind == "mid":
        intro = "👋 <b>Небольшой опрос на пару минут</b>\n\nПоможет понять, что в боте доработать в первую очередь."
    else:
        intro = "📄 <b>Финальный опрос по тестированию</b>\n\n5-7 вопросов, чтобы понять, куда двигать продукт дальше."
    intro_msg = await bot.send_message(uid, intro)
    # сохраним intro id в answers служебно
    _set_survey_state(uid, kind, answers={"_intro_msg": intro_msg.message_id})
    await _send_survey_step(bot, uid, kind, first, {})
    return True


def kb_start():
    return _kb([[BTN(text="Начать", callback_data="beta:start")]])


def kb_q1():
    return _single_inline("q1", [
        "Собственник 1–2 машин",
        "Собственник более 3 машин",
        "Логист / диспетчер",
        "Закупки / грузовладелец",
        "Руководитель отдела",
        "Сотрудник СБ",
        "Другое",
    ])


def kb_q2():
    return _single_inline("q2", [
        "Работаю один (1–2)",
        "Небольшой масштаб (3–10)",
        "Средний масштаб (11–50)",
        "Крупный масштаб (50+)",
        "Точно не могу оценить",
    ])


def kb_q3():
    return _single_inline("q3", [
        "Ежедневно",
        "Несколько раз в неделю",
        "Пару раз в месяц",
        "Редко",
    ])


def kb_q5():
    return _single_inline("q5", [
        "Да, были серьёзные случаи",
        "Были единичные/мелкие случаи",
        "Нет, проблем не было",
    ])


def kb_q7():
    return _single_inline(
        "q7",
        [
            "Активно тестировать",
            "Периодически пользоваться",
            "Не смогу участвовать",
        ],
    )


def kb_q8():
    return _single_inline(
        "q8",
        [
            "до 24",
            "25–34",
            "35–44",
            "45–54",
            "55+",
        ],
    )


def _single_inline(prefix: str, options: list[str]) -> InlineKeyboardMarkup:
    rows = [[InlineKeyboardButton(text=opt, callback_data=f"{prefix}:{idx}")] for idx, opt in enumerate(options)]
    return InlineKeyboardMarkup(inline_keyboard=rows)


async def _edit_or_send(target: Message, state: FSMContext, text: str, markup=None, *, force_new: bool = False) -> None:
    data = await state.get_data()
    last_msg_id = data.get("beta_last_msg_id")
    if force_new and last_msg_id:
        try:
            await target.bot.delete_message(chat_id=target.chat.id, message_id=last_msg_id)
        except Exception:
            pass
        last_msg_id = None
    try:
        if last_msg_id and not force_new:
            await target.bot.edit_message_text(
                chat_id=target.chat.id,
                message_id=last_msg_id,
                text=text,
                reply_markup=markup,
                parse_mode=ParseMode.HTML,
            )
            return
    except Exception:
        last_msg_id = None
    msg = await target.answer(text, reply_markup=markup)
    await state.update_data({"beta_last_msg_id": msg.message_id})


QUESTION_TEXTS: dict[str, dict[str, Any]] = {
    "q1": {"text": "1️⃣ <b>В какой роли вы работаете в логистике?</b>\nВыберите вариант, который точнее всего описывает вашу ситуацию.", "kb": kb_q1},
    "q2": {"text": "2️⃣ <b>Какой у вас масштаб работы?</b>\nОцените размер парка или команды в целом.", "kb": kb_q2},
    "q3": {"text": "3️⃣ <b>Как часто вы пользуетесь АТИ?</b>", "kb": kb_q3},
    "q4": {
        "text": (
            "4️⃣ <b>Что вы чаще всего делаете в АТИ?</b>\n"
            "Можно выбрать несколько пунктов, нажимая варианты. Когда закончите — нажмите «Готово»."
        ),
        "kb": lambda: build_q4_kb([]),
    },
    "q5": {
        "text": (
            "5️⃣ <b>Были ли у вас проблемы с контрагентами за последние 12 месяцев?</b>\n"
            "Примеры: задержка/отказ оплаты, исчезли, спорные документы, и т. п."
        ),
        "kb": kb_q5,
    },
    "q5_extra": {
        "text": (
            "Если готовы — опишите 1–2 ситуации коротко (по желанию).\n"
            "Можно в формате: «не оплатили», «пропали», «проблемы по документам».\n"
            "Если не хотите — нажмите «Пропустить»."
        ),
        "kb": lambda: _single_inline("q5extra", ["Пропустить"]),
    },
    "q6": {
        "text": (
            "6️⃣ <b>Как вы сейчас проверяете новых контрагентов?</b>\n"
            "Напишите коротко: что именно делаете перед началом работы."
        ),
        "kb": None,
    },
    "q7": {
        "text": (
            "7️⃣ <b>Готовность к участию в тестовом периоде</b>\n"
            "Он длится около 7 дней. Ожидаем, что вы выполните несколько проверок и оставите замечания/предложения."
        ),
        "kb": kb_q7,
    },
    "q8": {"text": "8️⃣ <b>Укажите ваш возраст</b>\nВыберите диапазон.", "kb": kb_q8},
    "q9": {
        "text": (
            "9️⃣ <b>Что для вас ключевое в таком боте?</b>\n"
            "1–2 предложения: какой результат вы хотите получать на практике."
        ),
        "kb": None,
    },
}


def next_step(current: str | None) -> str | None:
    if current is None:
        return STEP_ORDER[0]
    try:
        idx = STEP_ORDER.index(current)
    except ValueError:
        return None
    if idx + 1 < len(STEP_ORDER):
        return STEP_ORDER[idx + 1]
    return None


@router.message(CommandStart())
async def on_start(message: Message, state: FSMContext) -> None:
    await state.clear()
    if _is_admin(message):
        await message.answer(
            "Админ панель:\n"
            "/review — просмотреть заявки\n"
            "/send_approvals — разослать одобренным\n"
            "/send_rejects — разослать отказ\n"
            "/beta_stats — сводка по тестерам\n"
            "/beta_sleepers — список «спящих»\n"
            "/beta_export — выгрузка .xlsx\n"
            "/survey_mid — отправить mid-опрос\n"
            "/survey_final — отправить финальный опрос\n"
            "/survey_stats — сводка по опросам",
            reply_markup=None,
        )
        return
    msg = await message.answer(greeting(), reply_markup=kb_start())
    await state.update_data({"beta_last_msg_id": msg.message_id})


@router.callback_query(F.data == "beta:start")
async def on_start_flow(query: CallbackQuery, state: FSMContext) -> None:
    if _is_admin(query):
        await query.answer("Этот режим только для участников", show_alert=False)
        return
    await state.clear()
    await state.update_data(
        {"beta_step": None, "beta_answers": {}, "beta_last_msg_id": None, "beta_q4": [], "beta_return_to_confirm": False}
    )
    await state.update_data({"beta_last_msg_id": query.message.message_id})
    await _ask_step(query.message, state, "q1")
    await query.answer()


async def _ask_step(target: Message, state: FSMContext, step: str) -> None:
    data = QUESTION_TEXTS.get(step)
    if not data:
        return
    await state.update_data({"beta_step": step})
    if step == "q4":
        current = (await state.get_data()).get("beta_q4") or []
        await _edit_or_send(target, state, render_q4_text(current), build_q4_kb(current))
    else:
        kb_builder = data["kb"]
        text_steps = {"q5_extra", "q6", "q9"}
        force_new = step in text_steps or (await state.get_data()).get("beta_force_new_next", False)
        markup = kb_builder() if kb_builder else None
        if step in {"q6", "q9"}:
            markup = ForceReply(selective=False)
        await _edit_or_send(
            target,
            state,
            data["text"],
            markup,
            force_new=force_new,
        )
        if force_new:
            await state.update_data({"beta_force_new_next": False})


@router.callback_query(F.data.startswith("q1:") | F.data.startswith("q2:") | F.data.startswith("q3:") | F.data.startswith("q5:") | F.data.startswith("q7:") | F.data.startswith("q8:"))
async def on_single_choice(query: CallbackQuery, state: FSMContext) -> None:
    data = await state.get_data()
    step = data.get("beta_step")
    if not step or not query.data:
        await query.answer()
        return
    prefix, value = query.data.split(":", 1)
    if step != prefix:
        await query.answer()
        return
    answers = data.get("beta_answers", {})
    # восстановить текст из опций
    keyboard_map = {
        "q1": kb_q1,
        "q2": kb_q2,
        "q3": kb_q3,
        "q5": kb_q5,
        "q7": kb_q7,
        "q8": kb_q8,
    }
    opts = [btn.text for row in keyboard_map[prefix]().inline_keyboard for btn in row]
    try:
        idx = int(value)
        answers[prefix] = opts[idx]
    except Exception:
        answers[prefix] = opts[0] if opts else value
    await state.update_data({"beta_answers": answers})
    next_code = next_step(prefix)
    if prefix == "q5":
        if answers[prefix].lower().startswith("нет"):
            next_code = "q6"
        else:
            next_code = "q5_extra"
    if next_code == "confirm":
        await _ask_confirm(query.message, state)
    elif next_code:
        await _ask_step(query.message, state, next_code)
    else:
        await _finish(query.message, state)
    await query.answer()


@router.callback_query(F.data.startswith("q5extra:"))
async def on_q5extra_skip(query: CallbackQuery, state: FSMContext) -> None:
    data = await state.get_data()
    answers = data.get("beta_answers", {})
    answers["q5_extra"] = "Пропустить"
    await state.update_data({"beta_answers": answers})
    return_to_confirm = data.get("beta_return_to_confirm", False)
    if return_to_confirm:
        await state.update_data({"beta_return_to_confirm": False})
        await _ask_confirm(query.message, state)
    else:
        await _ask_step(query.message, state, "q6")
    await query.answer()


@router.message(F.text, ~F.text.regexp(r"^/"))
async def on_text_answers(message: Message, state: FSMContext) -> None:
    # если активен mid/final опрос — обрабатываем его и выходим
    if message.from_user:
        active = _active_survey(message.from_user.id)
        if active:
            await on_survey_text(message, state)
            return
    data = await state.get_data()
    step = data.get("beta_step")
    if step not in {"q5_extra", "q6", "q9"}:
        return
    text = (message.text or "").strip()
    answers = data.get("beta_answers", {})
    answers[step] = text
    await state.update_data({"beta_answers": answers})
    next_code = next_step(step)
    if step == "q9":
        next_code = "confirm"
    if data.get("beta_return_to_confirm", False):
        await state.update_data({"beta_return_to_confirm": False})
        await _ask_confirm(message, state, delete_previous=True)
        return
    if next_code == "confirm":
        await _ask_confirm(message, state, delete_previous=True)
    elif next_code:
        await state.update_data({"beta_force_new_next": True})
        await _ask_step(message, state, next_code)
    else:
        await _finish(message, state)


@router.callback_query(F.data.startswith("q4:"))
async def on_q4_toggle(query: CallbackQuery, state: FSMContext) -> None:
    data = await state.get_data()
    selected: list[str] = data.get("beta_q4") or []
    token = (query.data or "").split(":", 1)[-1]
    if token == "done":
        if not selected:
            await query.answer("Выберите хотя бы один вариант", show_alert=True)
            return
        answers = data.get("beta_answers", {})
        answers["q4"] = selected
        await state.update_data({"beta_answers": answers})
        return_to_confirm = data.get("beta_return_to_confirm", False)
        if return_to_confirm:
            await state.update_data({"beta_return_to_confirm": False})
            await _ask_confirm(query.message, state)
        else:
            await _ask_step(query.message, state, "q5")
        await query.answer()
        return

    # toggle
    if token in selected:
        selected = [x for x in selected if x != token]
    else:
        selected.append(token)
    await state.update_data({"beta_q4": selected})
    try:
        await query.message.edit_text(
            render_q4_text(selected),
            reply_markup=build_q4_kb(selected),
        )
    except TelegramBadRequest:
        await query.message.answer(render_q4_text(selected), reply_markup=build_q4_kb(selected))
    await query.answer()


async def _finish(message: Message, state: FSMContext, actor=None) -> None:
    answers = (await state.get_data()).get("beta_answers", {})
    # если q4 был в процессе редактирования, убедимся что сохранён
    if "q4" not in answers:
        selected = (await state.get_data()).get("beta_q4") or []
        if selected:
            answers["q4"] = selected
    sender = actor if actor is not None else message.from_user
    uid = sender.id if sender else 0
    username = sender.username if sender and sender.username else None
    # удалить сообщение подтверждения, чтобы оставить только финальное
    data = await state.get_data()
    last_msg_id = data.get("beta_last_msg_id")
    if last_msg_id:
        try:
            await message.bot.delete_message(chat_id=message.chat.id, message_id=last_msg_id)
        except Exception:
            pass
        await state.update_data({"beta_last_msg_id": None})
    record = {
        "uid": uid,
        "username": username,
        "answers": answers,
        "status": "pending",
        "notified": False,
        "granted": False,
        "granted_qty": 0,
    }
    rec_id = None
    try:
        rec_id = _insert_submission(record)
        record["id"] = rec_id
    except Exception:
        logger.exception("failed to persist beta form uid=%s", uid)
    try:
        if TARGET_CHAT:
            await message.bot.send_message(
                TARGET_CHAT,
                _format_submission(record),
                parse_mode=ParseMode.HTML,
            )
    except Exception:
        logger.exception("failed to push beta form to chat uid=%s", uid)
    await state.clear()
    await _edit_or_send(message, state, thanks(), markup=None, force_new=True)


async def _ask_confirm(message: Message, state: FSMContext, delete_previous: bool = False) -> None:
    answers = (await state.get_data()).get("beta_answers", {})
    summary = _pretty_summary(answers)
    if delete_previous:
        data = await state.get_data()
        last_msg_id = data.get("beta_last_msg_id")
        if last_msg_id:
            try:
                await message.bot.delete_message(chat_id=message.chat.id, message_id=last_msg_id)
            except Exception:
                pass
    inline_kb = InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="✅ Отправить", callback_data="confirm:send")],
            [InlineKeyboardButton(text="Изменить вопрос 1", callback_data="confirm:edit:1")],
            [
                InlineKeyboardButton(text="Вопрос 2", callback_data="confirm:edit:2"),
                InlineKeyboardButton(text="Вопрос 3", callback_data="confirm:edit:3"),
            ],
            [
                InlineKeyboardButton(text="Вопрос 4", callback_data="confirm:edit:4"),
                InlineKeyboardButton(text="Вопрос 5", callback_data="confirm:edit:5"),
            ],
            [
                InlineKeyboardButton(text="Вопрос 6", callback_data="confirm:edit:6"),
                InlineKeyboardButton(text="Вопрос 7", callback_data="confirm:edit:7"),
            ],
            [
                InlineKeyboardButton(text="Вопрос 8", callback_data="confirm:edit:8"),
                InlineKeyboardButton(text="Вопрос 9", callback_data="confirm:edit:9"),
            ],
        ]
    )
    await state.update_data({"beta_step": "confirm"})
    msg = await message.answer("Проверьте ответы:\n\n" + summary, reply_markup=inline_kb)
    await state.update_data({"beta_last_msg_id": msg.message_id})


@router.callback_query(ConfirmFilter("confirm:send"))
async def on_confirm_send(query: CallbackQuery, state: FSMContext) -> None:
    await _finish(query.message, state, actor=query.from_user)
    await query.answer()


@router.callback_query(ConfirmFilter("confirm:edit"))
async def on_confirm_edit(query: CallbackQuery, state: FSMContext) -> None:
    parts = (query.data or "").split(":")
    if len(parts) == 2:
        # Полный сброс
        await state.clear()
        await state.update_data({"beta_answers": {}, "beta_q4": [], "beta_return_to_confirm": True})
        await _ask_step(query.message, state, "q1")
        await query.answer("Измените ответы")
        return
    if len(parts) == 3:
        try:
            target_q = int(parts[2])
        except ValueError:
            await query.answer()
            return
        if 1 <= target_q <= 9:
            await state.update_data({"beta_step": None, "beta_return_to_confirm": True})
            await _ask_step(query.message, state, f"q{target_q}")
            await query.answer()
            return
    await query.answer()


def _is_admin(obj: Message | CallbackQuery) -> bool:
    uid = None
    if isinstance(obj, Message):
        uid = obj.from_user.id if obj.from_user else None
    else:
        uid = obj.from_user.id if obj.from_user else None
    return uid in ADMIN_IDS


# --- Storage on sqlite ----------------------------------------------------


def _get_conn() -> sqlite3.Connection:
    global db_conn
    if db_conn is None:
        DB_PATH.parent.mkdir(parents=True, exist_ok=True)
        db_conn = sqlite3.connect(DB_PATH, check_same_thread=False)
        db_conn.row_factory = sqlite3.Row
    return db_conn


def _init_db() -> None:
    conn = _get_conn()
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS submissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            uid INTEGER NOT NULL,
            username TEXT,
            answers TEXT NOT NULL,
            status TEXT DEFAULT 'pending',
            notified INTEGER DEFAULT 0,
            granted INTEGER DEFAULT 0,
            granted_qty INTEGER DEFAULT 0,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS beta_surveys (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            uid INTEGER NOT NULL,
            kind TEXT NOT NULL,
            status TEXT DEFAULT 'sent', -- sent | answered | skipped | expired
            answers TEXT,
            current_step TEXT,
            message_id INTEGER,
            sent_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            answered_at TIMESTAMP,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(uid, kind)
        )
        """
    )
    conn.commit()


def _insert_submission(record: dict[str, Any]) -> int:
    conn = _get_conn()
    cur = conn.execute(
        "INSERT INTO submissions (uid, username, answers, status, notified, granted, granted_qty) VALUES (?, ?, ?, ?, ?, ?, ?)",
        (
            record.get("uid"),
            record.get("username"),
            json.dumps(record.get("answers"), ensure_ascii=False),
            record.get("status", "pending"),
            1 if record.get("notified") else 0,
            1 if record.get("granted") else 0,
            record.get("granted_qty", 0),
        ),
    )
    conn.commit()
    return cur.lastrowid


def _update_submission(rec_id: int, **fields: Any) -> None:
    if not fields:
        return
    conn = _get_conn()
    sets = ", ".join([f"{k}=?" for k in fields.keys()])
    values = [json.dumps(v, ensure_ascii=False) if k == "answers" else v for k, v in fields.items()]
    values.append(rec_id)
    conn.execute(f"UPDATE submissions SET {sets}, updated_at=CURRENT_TIMESTAMP WHERE id=?", values)
    conn.commit()


def _row_to_dict(row: sqlite3.Row) -> dict[str, Any]:
    answers = row["answers"]
    try:
        answers = json.loads(answers)
    except Exception:
        answers = {}
    return {
        "id": row["id"],
        "uid": row["uid"],
        "username": row["username"],
        "answers": answers,
        "status": row["status"],
        "notified": bool(row["notified"]),
        "granted": bool(row["granted"]),
        "granted_qty": row["granted_qty"],
        "created_at": row["created_at"],
    }


def _get_next_pending(after_id: Optional[int] = None) -> Optional[dict[str, Any]]:
    conn = _get_conn()
    if after_id:
        cur = conn.execute(
            "SELECT * FROM submissions WHERE status='pending' AND id>? ORDER BY id ASC LIMIT 1", (after_id,)
        )
    else:
        cur = conn.execute("SELECT * FROM submissions WHERE status='pending' ORDER BY id ASC LIMIT 1")
    row = cur.fetchone()
    return _row_to_dict(row) if row else None


def _get_all_by_status(status: str, notified: Optional[bool] = None) -> list[dict[str, Any]]:
    conn = _get_conn()
    if notified is None:
        cur = conn.execute("SELECT * FROM submissions WHERE status=? ORDER BY id ASC", (status,))
    else:
        cur = conn.execute(
            "SELECT * FROM submissions WHERE status=? AND notified=? ORDER BY id ASC",
            (status, 1 if notified else 0),
        )
    return [_row_to_dict(r) for r in cur.fetchall()]


def _set_notified(rec_id: int) -> None:
    _update_submission(rec_id, notified=1)


def _set_status(rec_id: int, status: str) -> None:
    _update_submission(rec_id, status=status)


def _set_granted(rec_id: int, qty: int) -> None:
    _update_submission(rec_id, granted=1, granted_qty=qty)


def _approved_submissions() -> list[dict[str, Any]]:
    conn = _get_conn()
    cur = conn.execute("SELECT uid, username FROM submissions WHERE status='approved'")
    res = []
    for row in cur.fetchall():
        uid = row["uid"]
        if uid is None:
            continue
        res.append({"uid": int(uid), "username": row["username"]})
    return res


def _approved_uids() -> list[int]:
    return [rec["uid"] for rec in _approved_submissions()]


def _user_link(uid: int, username: Optional[str]) -> str:
    link = f"tg://user?id={uid}"
    text = f"@{username}" if username else str(uid)
    return f'<a href="{link}">{text}</a>'


def _as_dt(val):
    if val is None:
        return None
    dt = None
    if isinstance(val, datetime):
        dt = val
    elif isinstance(val, str):
        try:
            dt = datetime.fromisoformat(val)
        except Exception:
            try:
                if val.endswith("Z"):
                    dt = datetime.fromisoformat(val.replace("Z", "+00:00"))
            except Exception:
                dt = None
    if dt is None:
        return None
    if dt.tzinfo is not None:
        dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
    return dt


# --- Surveys storage ------------------------------------------------------


def _get_survey(uid: int, kind: str, statuses: Optional[set[str]] = None) -> Optional[dict[str, Any]]:
    conn = _get_conn()
    if statuses:
        placeholders = ",".join(["?"] * len(statuses))
        params = list(statuses) + [uid, kind]
        cur = conn.execute(
            f"SELECT * FROM beta_surveys WHERE status IN ({placeholders}) AND uid=? AND kind=? LIMIT 1",
            params,
        )
    else:
        cur = conn.execute("SELECT * FROM beta_surveys WHERE uid=? AND kind=? LIMIT 1", (uid, kind))
    row = cur.fetchone()
    if not row:
        return None
    answers = row["answers"]
    try:
        answers = json.loads(answers) if answers else {}
    except Exception:
        answers = {}
    return {
        "id": row["id"],
        "uid": row["uid"],
        "kind": row["kind"],
        "status": row["status"],
        "answers": answers,
        "current_step": row["current_step"],
        "message_id": row["message_id"],
        "sent_at": row["sent_at"],
        "answered_at": row["answered_at"],
    }


def _create_survey(uid: int, kind: str, status: str = "sent", step: str | None = None, message_id: int | None = None) -> int:
    conn = _get_conn()
    cur = conn.execute(
        """
        INSERT INTO beta_surveys (uid, kind, status, current_step, message_id)
        VALUES (?, ?, ?, ?, ?)
        ON CONFLICT(uid, kind) DO UPDATE SET status=excluded.status, current_step=excluded.current_step, message_id=excluded.message_id, sent_at=CURRENT_TIMESTAMP, updated_at=CURRENT_TIMESTAMP
        """,
        (uid, kind, status, step, message_id),
    )
    conn.commit()
    return cur.lastrowid or conn.execute("SELECT id FROM beta_surveys WHERE uid=? AND kind=?", (uid, kind)).fetchone()[0]


def _set_survey_state(uid: int, kind: str, **fields: Any) -> None:
    if not fields:
        return
    conn = _get_conn()
    sets = ", ".join([f"{k}=?" for k in fields.keys()])
    values = [json.dumps(v, ensure_ascii=False) if k == "answers" else v for k, v in fields.items()]
    values.extend([uid, kind])
    conn.execute(
        f"UPDATE beta_surveys SET {sets}, updated_at=CURRENT_TIMESTAMP WHERE uid=? AND kind=?",
        values,
    )
    conn.commit()


def _active_survey(uid: int) -> Optional[dict[str, Any]]:
    return _get_survey(uid, kind="mid", statuses={"sent"}) or _get_survey(uid, kind="final", statuses={"sent"})


async def main() -> None:
    global quota_service
    _init_db()
    await dal.init_db()
    init_activity_db()
    quota_service = QuotaService(tz=cfg.tz)
    bot = Bot(token=BOT_TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    dp = Dispatcher()
    dp.include_router(router)
    asyncio.create_task(reminders_loop(bot))
    used = set(dp.resolve_used_update_types() or [])
    used.add("chat_join_request")
    await dp.start_polling(bot, allowed_updates=used)


def _format_submission(record: dict[str, Any]) -> str:
    uid = record.get("uid")
    username = record.get("username")
    rec_id = record.get("id")
    answers: dict[str, Any] = record.get("answers") or {}
    lines = [
        "<b>Заявка в бету</b>",
        f"id: <code>{rec_id}</code>",
        f"uid: <code>{uid}</code>",
        f"username: @{username}" if username else "username: —",
        "",
        "<b>Ответы:</b>",
        f"1) Роль: {answers.get('q1', '—')}",
        f"2) Масштаб: {answers.get('q2', '—')}",
        f"3) Частота АТИ: {answers.get('q3', '—')}",
        f"4) Действия в АТИ: {', '.join(answers.get('q4', [])) if isinstance(answers.get('q4'), list) else answers.get('q4', '—')}",
        f"5) Проблемы за год: {answers.get('q5', '—')}",
        f"5b) Детали проблем: {answers.get('q5_extra', '—')}" if answers.get("q5_extra") not in (None, "", "Пропустить") else None,
        f"6) Как проверяет: {answers.get('q6', '—')}",
        f"7) Готовность: {answers.get('q7', '—')}",
        f"8) Возраст: {answers.get('q8', '—')}",
        f"9) Ожидания: {answers.get('q9', '—')}",
    ]
    return "\n".join([line for line in lines if line is not None])


def _pretty_summary(answers: dict[str, Any]) -> str:
    lines = [
        f"1) Роль: {answers.get('q1', '—')}",
        f"2) Масштаб: {answers.get('q2', '—')}",
        f"3) Частота АТИ: {answers.get('q3', '—')}",
        f"4) Действия в АТИ: {', '.join(answers.get('q4', [])) if isinstance(answers.get('q4'), list) else answers.get('q4', '—')}",
        f"5) Проблемы за год: {answers.get('q5', '—')}",
        f"5b) Детали проблем: {answers.get('q5_extra', '—')}" if answers.get("q5_extra") not in (None, "", "Пропустить") else None,
        f"6) Как проверяет: {answers.get('q6', '—')}",
        f"7) Готовность: {answers.get('q7', '—')}",
        f"8) Возраст: {answers.get('q8', '—')}",
        f"9) Ожидания: {answers.get('q9', '—')}",
    ]
    return "\n".join([line for line in lines if line is not None])


def _format_for_admin(rec: dict[str, Any]) -> str:
    uid = rec.get("uid")
    username = rec.get("username")
    status = rec.get("status", "pending")
    rec_id = rec.get("id")
    answers: dict[str, Any] = rec.get("answers") or {}
    lines = [
        f"<b>Заявка #{rec_id}</b> (status: {status})",
        f"uid: <code>{uid}</code>",
        f"username: @{username}" if username else "username: —",
        "",
        _pretty_summary(answers),
    ]
    return "\n".join(lines)


# --- Reminder scheduler ----------------------------------------------------


REMINDER_SCAN_SEC = int(os.getenv("REMINDER_SCAN_SEC", "600") or 600)
REMINDER_24_MAX = int(os.getenv("REMINDER_24_MAX", "3") or 3)


async def _maybe_send_surveys(bot: Bot, activity_rows: list[sqlite3.Row], now: datetime) -> None:
    # авто-отправка mid/final по времени
    by_uid = {row["uid"]: row for row in activity_rows}
    approved = _approved_uids()
    sent_mid = []
    sent_final = []
    for uid in approved:
        row = by_uid.get(uid)
        if not row:
            continue
        first_start = row["first_start_at"]
        first_dt = None
        if isinstance(first_start, str):
            try:
                first_dt = datetime.fromisoformat(first_start)
            except Exception:
                logger.warning("cannot parse first_start_at for uid=%s raw=%s", uid, first_start)
                first_dt = None
        elif isinstance(first_start, datetime):
            first_dt = first_start
        if first_dt is None:
            logger.debug("skip survey scheduling: no first_dt uid=%s", uid)
            continue
        if first_dt.tzinfo is None:
            logger.warning("first_start_at is naive, assuming UTC uid=%s raw=%s", uid, first_start)
            first_dt = first_dt.replace(tzinfo=timezone.utc)
        # mid
        try:
            delta_sec = (now - first_dt).total_seconds()
        except Exception:
            logger.exception(
                "failed to compute delta for surveys uid=%s now=%s first_dt=%s",
                uid,
                now,
                first_dt,
            )
            continue
        if not _get_survey(uid, "mid", statuses={"sent", "answered"}) and delta_sec >= SURVEY_MID_DELAY_HOURS * 3600:
            logger.info("auto-sending mid survey uid=%s delta_hours=%.2f", uid, delta_sec / 3600)
            await _start_survey(bot, uid, "mid")
            sent_mid.append(uid)
            continue  # чтобы не наваливать сразу финальный
        # final
        if not _get_survey(uid, "final", statuses={"sent", "answered"}) and delta_sec >= SURVEY_FINAL_DELAY_HOURS * 3600:
            logger.info("auto-sending final survey uid=%s delta_hours=%.2f", uid, delta_sec / 3600)
            await _start_survey(bot, uid, "final")
            sent_final.append(uid)
    if sent_mid or sent_final:
        await _notify_admins(
            bot,
            f"Опросы разосланы автоматически: mid={len(sent_mid)}, final={len(sent_final)}.",
        )


async def reminders_loop(bot: Bot) -> None:
    while True:
        try:
            await _run_reminders(bot)
        except Exception:
            logger.exception("reminder scan failed")
        await asyncio.sleep(REMINDER_SCAN_SEC)


async def _run_reminders(bot: Bot) -> None:
    now = now_utc()
    rows = get_activity_rows()
    await _maybe_send_surveys(bot, rows, now)

    def _as_dt(val):
        if val is None:
            return None
        if isinstance(val, datetime):
            if val.tzinfo is None:
                logger.warning("reminder datetime is naive, assuming UTC val=%s", val)
                return val.replace(tzinfo=timezone.utc)
            return val
        if isinstance(val, str):
            try:
                dt = datetime.fromisoformat(val)
                if dt.tzinfo is None:
                    logger.warning("reminder datetime string is naive, assuming UTC raw=%s", val)
                    dt = dt.replace(tzinfo=timezone.utc)
                return dt
            except Exception:
                logger.warning("failed to parse datetime string raw=%s", val)
                return None
        return None

    for row in rows:
        uid = row["uid"]
        if row["inactive"]:
            continue
        if _active_survey(uid):
            # не отвлекаем напоминаниями, если пользователь в процессе опроса
            continue
        first_start = _as_dt(row["first_start_at"])
        reminder_start_sent = bool(row["reminder_start_sent"])
        reminder_3h_sent = bool(row["reminder_3h_sent"])
        reminder_24_sent_at = _as_dt(row["reminder_24_sent_at"])
        reminder_24_count = row["reminder_24_count"] or 0
        reminder_48_sent_at = _as_dt(row["reminder_48_sent_at"])

        total_checks, first_check, last_check = await get_history_stats(uid)

        # Стартовое сообщение — один раз
        if first_start and not reminder_start_sent:
            await bot.send_message(uid, reminder_start_text())
            mark_start_reminder_sent(uid)

        # 3 часа без проверок
        if first_start and not reminder_3h_sent and total_checks == 0:
            try:
                delta_minutes = (now - first_start).total_seconds() / 60
            except Exception:
                logger.exception("delta calc failed (3h) uid=%s now=%s first=%s", uid, now, first_start)
                delta_minutes = None
            if delta_minutes >= 180:
                kb = InlineKeyboardMarkup(
                    inline_keyboard=[
                        [InlineKeyboardButton(text="🔎 Сделать проверку", url="https://t.me/antifraudbetabot?start=")]
                    ]
                )
                await bot.send_message(uid, reminder_3h_text(), reply_markup=kb)
                mark_3h_sent(uid)

        # 24 часа без новых проверок (шлём даже если проверок не было)
        anchor_24 = last_check or first_start
        if anchor_24 and reminder_24_count < REMINDER_24_MAX:
            try:
                delta_minutes = (now - anchor_24).total_seconds() / 60
            except Exception:
                logger.exception("delta calc failed (24h) uid=%s now=%s anchor=%s", uid, now, anchor_24)
                delta_minutes = None
            if delta_minutes >= 1440:
                if reminder_24_sent_at is None or (now - reminder_24_sent_at).total_seconds() >= 1440 * 60:
                    kb = InlineKeyboardMarkup(
                        inline_keyboard=[
                            [InlineKeyboardButton(text="🔎 Сделать проверку", url="https://t.me/antifraudbetabot?start=")]
                        ]
                    )
                    await bot.send_message(uid, reminder_24h_text(), reply_markup=kb)
                    mark_24h_sent(uid)

        # 48 часов без активности (шлём даже если проверок не было)
        anchor = last_check or first_start
        if anchor:
            try:
                delta_minutes = (now - anchor).total_seconds() / 60
            except Exception:
                logger.exception("delta calc failed (48h) uid=%s now=%s anchor=%s", uid, now, anchor)
                delta_minutes = None
            if delta_minutes >= 2880:
                if reminder_48_sent_at is None or (now - reminder_48_sent_at).total_seconds() >= 2880 * 60:
                    kb = InlineKeyboardMarkup(
                        inline_keyboard=[
                            [InlineKeyboardButton(text="🔎 Сделать проверку", url="https://t.me/antifraudbetabot?start=")],
                            [InlineKeyboardButton(text="🚫 Не напоминать", callback_data=f"rem:off:{uid}")],
                        ]
                    )
                    await bot.send_message(uid, reminder_48h_text(), reply_markup=kb)
                    mark_48h_sent(uid)


@router.callback_query(F.data.startswith("rem:off:"))
async def on_reminder_off(query: CallbackQuery) -> None:
    token = query.data.split(":")[-1]
    try:
        uid = int(token)
    except ValueError:
        await query.answer()
        return
    mark_inactive(uid)
    await query.answer("Ок, больше не будем напоминать", show_alert=False)


# --- Survey handlers ------------------------------------------------------


def _survey_expected(uid: int) -> Optional[dict[str, Any]]:
    s = _active_survey(uid)
    return s


def _survey_store_answer(uid: int, kind: str, code: str, value: Any, base_answers: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    survey = _get_survey(uid, kind, statuses={"sent"})
    answers = base_answers if base_answers is not None else (survey["answers"] if survey else {})
    answers[code] = value
    _set_survey_state(uid, kind, answers=answers)
    return answers


async def _advance_survey(bot: Bot, uid: int, kind: str, current: str, answers: dict[str, Any]) -> None:
    next_code = _survey_next(kind, current, answers)
    if next_code:
        await _send_survey_step(bot, uid, kind, next_code, answers)
        return
    _set_survey_state(
        uid,
        kind,
        status="answered",
        answered_at=datetime.now(timezone.utc).replace(tzinfo=None),
        current_step=None,
    )
    await bot.send_message(uid, "✅ Спасибо, ответы сохранены.")


@router.callback_query(F.data.startswith("sv:"))
async def on_survey_choice(query: CallbackQuery) -> None:
    parts = (query.data or "").split(":")
    if len(parts) != 4:
        await query.answer()
        return
    _, kind, code, idx_raw = parts
    uid = query.from_user.id
    survey = _get_survey(uid, kind, statuses={"sent"})
    if not survey or survey.get("current_step") != code:
        await query.answer()
        return
    # удалить предыдущее сообщение с вопросом и интро, если ещё не удалено
    answers = survey.get("answers") or {}
    intro_id = answers.pop("_intro_msg", None)
    try:
        msg_id = survey.get("message_id")
        if msg_id:
            await query.bot.delete_message(query.message.chat.id, msg_id)
    except Exception:
        pass
    if intro_id:
        try:
            await query.bot.delete_message(query.message.chat.id, intro_id)
        except Exception:
            pass
    step = _survey_step(kind, code)
    if not step:
        await query.answer()
        return
    opts = step.get("options") or []
    if step["type"] == "scale5":
        opts = step.get("labels") or ["1", "2", "3", "4", "5"]
    try:
        idx = int(idx_raw)
        value = opts[idx] if idx < len(opts) else opts[0]
    except Exception:
        value = opts[0] if opts else idx_raw
    answers = _survey_store_answer(uid, kind, code, value, base_answers=answers)
    await _advance_survey(query.bot, uid, kind, code, answers)
    await query.answer()


@router.callback_query(F.data.startswith("svskip:"))
async def on_survey_skip(query: CallbackQuery) -> None:
    parts = (query.data or "").split(":")
    if len(parts) != 3:
        await query.answer()
        return
    _, kind, code = parts
    uid = query.from_user.id
    survey = _get_survey(uid, kind, statuses={"sent"})
    if not survey:
        await query.answer()
        return
    answers = survey.get("answers") or {}
    intro_id = answers.pop("_intro_msg", None)
    try:
        msg_id = survey.get("message_id")
        if msg_id:
            await query.bot.delete_message(query.message.chat.id, msg_id)
    except Exception:
        pass
    if intro_id:
        try:
            await query.bot.delete_message(query.message.chat.id, intro_id)
        except Exception:
            pass
    answers = _survey_store_answer(uid, kind, code, "Пропустить", base_answers=answers)
    await _advance_survey(query.bot, uid, kind, code, answers)
    await query.answer("Пропущено")


@router.message(F.text, ~F.text.regexp(r"^/"))
async def on_survey_text(message: Message, state: FSMContext) -> None:
    uid = message.from_user.id if message.from_user else None
    if uid is None:
        return
    survey = _active_survey(uid)
    if not survey:
        # fall back to основная анкета
        await on_text_answers(message, state)
        return
    current = survey.get("current_step")
    step = _survey_step(survey["kind"], current) if current else None
    if not step or step["type"] not in {"text", "text_optional"}:
        return
    answers = survey.get("answers") or {}
    intro_id = answers.pop("_intro_msg", None)
    try:
        msg_id = survey.get("message_id")
        if msg_id:
            await message.bot.delete_message(message.chat.id, msg_id)
    except Exception:
        pass
    if intro_id:
        try:
            await message.bot.delete_message(message.chat.id, intro_id)
        except Exception:
            pass
    text = (message.text or "").strip()
    answers = _survey_store_answer(uid, survey["kind"], current, text, base_answers=answers)
    await _advance_survey(message.bot, uid, survey["kind"], current, answers)


async def _agg_checks(uids: list[int], since: Optional[datetime] = None) -> dict[int, dict[str, Any]]:
    if not uids:
        return {}
    async with dal.Session() as session:
        stmt = select(
            dal.history.c.uid,
            func.count(dal.history.c.id).label("cnt"),
            func.max(dal.history.c.ts).label("last_ts"),
        ).where(dal.history.c.uid.in_(uids))
        if since is not None:
            stmt = stmt.where(dal.history.c.ts >= since)
        stmt = stmt.group_by(dal.history.c.uid)
        result = await session.execute(stmt)
        data = {}
        for row in result:
            data[int(row.uid)] = {"cnt": int(row.cnt or 0), "last_ts": row.last_ts}
        return data


def _start_of_today_msk() -> datetime:
    tz = timezone(timedelta(hours=3))
    today = datetime.now(tz).date()
    start_local = datetime.combine(today, time(0, 0), tzinfo=tz)
    return start_local.astimezone(timezone.utc)


@router.message(Command("beta_stats"))
async def on_beta_stats(message: Message) -> None:
    if not _is_admin(message):
        await message.answer("Нет доступа")
        return
    subs = _approved_submissions()
    uids = [rec["uid"] for rec in subs]
    uname_map = {rec["uid"]: rec.get("username") for rec in subs}
    if not uids:
        await message.answer("Нет одобренных тестеров.")
        return
    stats_all = await _agg_checks(uids)
    stats_today = await _agg_checks(uids, since=_start_of_today_msk())

    total = len(uids)
    checks_map = {uid: info["cnt"] for uid, info in stats_all.items()}
    active1 = sum(1 for uid in uids if checks_map.get(uid, 0) >= 1)
    active5 = sum(1 for uid in uids if checks_map.get(uid, 0) >= 5)
    zero = sum(1 for uid in uids if checks_map.get(uid, 0) == 0)
    avg_checks = (sum(checks_map.values()) / active1) if active1 > 0 else 0

    today_total = sum(stats_today[uid]["cnt"] for uid in stats_today)

    top = sorted(checks_map.items(), key=lambda x: x[1], reverse=True)[:10]
    top_lines = []
    for uid, cnt in top:
        top_lines.append(f"{_user_link(uid, uname_map.get(uid))} (uid {uid}): {cnt}")

    text = (
        "<b>Статистика беты</b>\n"
        f"Всего одобрено: {total}\n"
        f"Сделали ≥1 проверку: {active1}\n"
        f"Сделали ≥5 проверок: {active5}\n"
        f"Не сделали ни одной: {zero}\n"
        f"Среднее проверок на активного: {avg_checks:.2f}\n"
        f"Проверок сегодня (MSK): {today_total}\n\n"
        "<b>Топ-10 по проверкам</b>\n"
        + ("\n".join(top_lines) if top_lines else "—")
    )
    await message.answer(text)


@router.message(Command("beta_sleepers"))
async def on_beta_sleepers(message: Message) -> None:
    if not _is_admin(message):
        await message.answer("Нет доступа")
        return
    subs = _approved_submissions()
    uids = [rec["uid"] for rec in subs]
    uname_map = {rec["uid"]: rec.get("username") for rec in subs}
    if not uids:
        await message.answer("Нет одобренных тестеров.")
        return
    stats_all = await _agg_checks(uids)
    now = now_utc()
    sleepers = []
    for uid in uids:
        info = stats_all.get(uid, {"cnt": 0, "last_ts": None})
        cnt = info["cnt"]
        last_ts = info["last_ts"]
        if cnt == 0:
            sleepers.append((uid, cnt, last_ts))
        else:
            if last_ts and (now - last_ts).total_seconds() >= 48 * 3600:
                sleepers.append((uid, cnt, last_ts))
    if not sleepers:
        await message.answer("Нет спящих тестеров.")
        return
    lines = []
    for uid, cnt, last_ts in sleepers:
        ts_text = last_ts.isoformat() if last_ts else "—"
        lines.append(f"{_user_link(uid, uname_map.get(uid))} (uid {uid}), checks={cnt}, last={ts_text}")
    await message.answer("<b>Слиперы</b>\n" + "\n".join(lines))


async def _bulk_send_survey(kind: str, message: Message) -> None:
    if not _is_admin(message):
        await message.answer("Нет доступа")
        return
    rows = get_activity_rows()
    now = now_utc()
    sent = 0
    for row in rows:
        uid = row["uid"]
        if uid not in _approved_uids():
            continue
        if _get_survey(uid, kind, statuses={"sent", "answered"}):
            continue
        first_start = row["first_start_at"]
        first_dt = None
        if isinstance(first_start, str):
            try:
                first_dt = datetime.fromisoformat(first_start)
            except Exception:
                first_dt = None
        elif isinstance(first_start, datetime):
            first_dt = first_start
        if first_dt is None:
            continue
        delay_h = SURVEY_MID_DELAY_HOURS if kind == "mid" else SURVEY_FINAL_DELAY_HOURS
        if (now - first_dt).total_seconds() >= delay_h * 3600:
            await _start_survey(message.bot, uid, kind)
            sent += 1
    await message.answer(f"Опрос {kind}: отправлено {sent}")


@router.message(Command("survey_mid"))
async def cmd_survey_mid(message: Message) -> None:
    await _bulk_send_survey("mid", message)


@router.message(Command("survey_final"))
async def cmd_survey_final(message: Message) -> None:
    await _bulk_send_survey("final", message)


@router.message(Command("survey_user"))
async def cmd_survey_user(message: Message) -> None:
    if not _is_admin(message):
        await message.answer("Нет доступа")
        return
    parts = (message.text or "").split()
    if len(parts) != 3:
        await message.answer("Формат: /survey_user mid|final <uid>")
        return
    _, kind, uid_raw = parts
    if kind not in {"mid", "final"}:
        await message.answer("Укажите mid или final")
        return
    try:
        uid = int(uid_raw)
    except ValueError:
        await message.answer("uid должен быть числом")
        return
    started = await _start_survey(message.bot, uid, kind)
    if started:
        await message.answer(f"Опрос {kind} отправлен пользователю {uid}.")
    else:
        await message.answer("Не удалось отправить опрос (возможно, уже пройден).")


@router.message(Command("survey_stats"))
async def cmd_survey_stats(message: Message) -> None:
    if not _is_admin(message):
        await message.answer("Нет доступа")
        return
    conn = _get_conn()
    cur = conn.execute(
        """
        SELECT kind, status, COUNT(*) as cnt
        FROM beta_surveys
        GROUP BY kind, status
        """
    )
    rows = cur.fetchall()
    if not rows:
        await message.answer("Нет данных по опросам.")
        return
    summary = {}
    for row in rows:
        summary.setdefault(row["kind"], {})[row["status"]] = row["cnt"]
    lines = []
    for kind in ("mid", "final"):
        stats = summary.get(kind, {})
        lines.append(
            f"{kind}: sent={stats.get('sent', 0)}, answered={stats.get('answered', 0)}, skipped={stats.get('skipped', 0)}, expired={stats.get('expired', 0)}"
        )
    await message.answer("\n".join(lines))


@router.message(Command("beta_export"))
async def on_beta_export(message: Message) -> None:
    if not _is_admin(message):
        await message.answer("Нет доступа")
        return
    if Workbook is None:
        await message.answer("Для экспорта нужен openpyxl. Установите пакет и перезапустите бота.")
        return
    subs = _approved_submissions()
    if not subs:
        await message.answer("Нет данных для экспорта.")
        return
    uids = [rec["uid"] for rec in subs]
    uname_map = {rec["uid"]: rec.get("username") for rec in subs}
    stats_all = await _agg_checks(uids)
    stats_today = await _agg_checks(uids, since=_start_of_today_msk())
    # convert rows to dict to avoid Row.get issues
    activity_map = {}
    for row in get_activity_rows():
        activity_map[row["uid"]] = {k: row[k] for k in row.keys()}

    wb = Workbook()
    ws = wb.active
    ws.title = "beta_stats"
    headers = [
        "uid",
        "username",
        "tg_link",
        "total_checks",
        "checks_today",
        "first_check",
        "last_check",
        "first_start_at",
        "reminder_start_sent",
        "reminder_3h_sent",
        "reminder_24_sent_at",
        "reminder_24_count",
        "reminder_48_sent_at",
        "inactive",
    ]
    ws.append(headers)

    for uid in uids:
        total = stats_all.get(uid, {}).get("cnt", 0)
        today = stats_today.get(uid, {}).get("cnt", 0)
        _, first_check, last_check = await get_history_stats(uid)
        act = activity_map.get(uid, {})
        ws.append(
            [
                uid,
                uname_map.get(uid),
                f"https://t.me/{uname_map.get(uid)}" if uname_map.get(uid) else f"tg://user?id={uid}",
                total,
                today,
                _as_dt(first_check),
                _as_dt(last_check),
                _as_dt(act.get("first_start_at")),
                bool(act.get("reminder_start_sent")),
                bool(act.get("reminder_3h_sent")),
                _as_dt(act.get("reminder_24_sent_at")),
                act.get("reminder_24_count"),
                _as_dt(act.get("reminder_48_sent_at")),
                bool(act.get("inactive")),
            ]
        )

    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        wb.save(tmp.name)
        wb.close()
        tmp_path = tmp.name
    try:
        await message.answer_document(FSInputFile(tmp_path, filename="beta_stats.xlsx"))
    finally:
        try:
            os.remove(tmp_path)
        except Exception:
            pass


def _kb_review(rec_id: int, next_id: Optional[int]) -> InlineKeyboardMarkup:
    return _kb(
        [
            [
                BTN(text="✅ Одобрить", callback_data=f"adm:approve:{rec_id}"),
                BTN(text="❌ Отклонить", callback_data=f"adm:reject:{rec_id}"),
            ],
            [BTN(text="➡️ Далее", callback_data=f"adm:skip:{next_id or 0}")],
        ]
    )


def _kb_empty() -> InlineKeyboardMarkup:
    return _kb([[BTN(text="Обновить", callback_data="adm:refresh")]])


async def _send_next_pending(target: Message, after_id: Optional[int] = None) -> None:
    logger.info("send_next_pending after_id=%s", after_id)
    rec = _get_next_pending(after_id)
    if not rec:
        logger.info("queue empty")
        await target.answer("Очередь пуста", reply_markup=_kb_empty())
        return
    logger.info("showing rec id=%s status=%s", rec.get("id"), rec.get("status"))
    # найти следующий pending id для кнопки skip
    next_rec = _get_next_pending(rec.get("id"))
    next_id = next_rec.get("id") if next_rec else None
    await target.answer(_format_for_admin(rec), reply_markup=_kb_review(rec.get("id"), next_id))


@router.message(Command("review"))
async def on_review(message: Message, state: FSMContext) -> None:
    if not _is_admin(message):
        await message.answer("Нет доступа")
        return
    logger.info("admin %s requested /review", message.from_user.id if message.from_user else None)
    try:
        await _send_next_pending(message)
    except Exception:
        logger.exception("failed to send next pending")
        await message.answer("Ошибка при загрузке очереди")


@router.callback_query(F.data == "adm:refresh")
async def on_refresh(query: CallbackQuery, state: FSMContext) -> None:
    if not _is_admin(query):
        await query.answer("Нет доступа", show_alert=False)
        return
    await _send_next_pending(query.message)
    await query.answer()


@router.callback_query(F.data.startswith("adm:skip:"))
async def on_skip(query: CallbackQuery, state: FSMContext) -> None:
    if not _is_admin(query):
        await query.answer("Нет доступа", show_alert=False)
        return
    token = query.data.split(":")[-1]
    try:
        after_id = int(token) if token not in ("", "0") else None
    except ValueError:
        after_id = None
    await _send_next_pending(query.message, after_id=after_id)
    await query.answer()


@router.callback_query(F.data.startswith("adm:approve:"))
async def on_admin_approve(query: CallbackQuery, state: FSMContext) -> None:
    if not _is_admin(query):
        await query.answer("Нет доступа", show_alert=False)
        return
    parts = query.data.split(":")
    if len(parts) < 3:
        await query.answer()
        return
    try:
        rec_id = int(parts[2])
    except ValueError:
        await query.answer()
        return
    _set_status(rec_id, "approved")
    await _send_next_pending(query.message, after_id=rec_id)
    await query.answer("Одобрено")


@router.callback_query(F.data.startswith("adm:reject:"))
async def on_admin_reject(query: CallbackQuery, state: FSMContext) -> None:
    if not _is_admin(query):
        await query.answer("Нет доступа", show_alert=False)
        return
    parts = query.data.split(":")
    if len(parts) < 3:
        await query.answer()
        return
    try:
        rec_id = int(parts[2])
    except ValueError:
        await query.answer()
        return
    _set_status(rec_id, "rejected")
    await _send_next_pending(query.message, after_id=rec_id)
    await query.answer("Отклонено")


@router.message(Command("send_approvals"))
async def on_send_approvals(message: Message, state: FSMContext) -> None:
    if not _is_admin(message):
        await message.answer("Нет доступа")
        return
    data = _get_all_by_status("approved", notified=False)
    count = 0
    for rec in data:
        uid = rec.get("uid")
        if not uid:
            continue
        try:
            # гарантируем наличие пользователя в основной БД, чтобы сработал FK
            try:
                await dal.ensure_user(uid, rec.get("username"), None, None)
            except Exception:
                logger.exception("ensure_user failed for uid=%s", uid)
            ensure_activity(uid, now_utc())
            if quota_service and GRANT_QTY > 0 and not rec.get("granted"):
                try:
                    await quota_service.add(uid, GRANT_QTY, source="beta-auto")
                    _set_granted(rec["id"], GRANT_QTY)
                except IntegrityError:
                    logger.warning("skip quota grant: user %s not in users table", uid)
            # не пытаемся писать ботам
            if str(uid).endswith("bot") or (rec.get("username") and str(rec.get("username")).lower().endswith("bot")):
                logger.warning("skip notify for bot uid=%s", uid)
                _set_notified(rec["id"])
                continue
            add_to_whitelist(uid)
            if GROUP_LINK:
                kb = InlineKeyboardMarkup(
                    inline_keyboard=[[InlineKeyboardButton(text="💬 Вступить в чат участников", url=GROUP_LINK)]]
                )
                await message.bot.send_message(uid, accept_text(""), reply_markup=kb)
            _set_notified(rec["id"])
            count += 1
        except Exception:
            logger.exception("failed to notify approval uid=%s", uid)
    await message.answer(f"Отправлено уведомлений: {count}")


@router.message(Command("send_rejects"))
async def on_send_rejects(message: Message, state: FSMContext) -> None:
    if not _is_admin(message):
        await message.answer("Нет доступа")
        return
    data = _get_all_by_status("rejected", notified=False)
    count = 0
    for rec in data:
        uid = rec.get("uid")
        if not uid:
            continue
        try:
            await message.bot.send_message(uid, decline_text())
            _set_notified(rec["id"])
            count += 1
        except Exception:
            logger.exception("failed to notify reject uid=%s", uid)
    await message.answer(f"Отправлено уведомлений (отказ): {count}")


@router.chat_join_request()
async def on_join_request(event: ChatJoinRequest) -> None:
    # если задан конкретный чат для проверки, фильтруем по нему
    if GROUP_ID is not None and event.chat.id != GROUP_ID:
        return
    uid = event.from_user.id
    try:
        if is_whitelisted(uid):
            await event.bot.approve_chat_join_request(chat_id=event.chat.id, user_id=uid)
            logger.info("approved join request uid=%s chat=%s", uid, event.chat.id)
            ensure_activity(uid, now_utc())
            if TEST_BOT_LINK:
                bot_name = TEST_BOT_LINK.lstrip("@")
                link = TEST_BOT_LINK
                if not link.lower().startswith("http"):
                    link = f"https://t.me/{bot_name}?start="
                kb = InlineKeyboardMarkup(
                    inline_keyboard=[[InlineKeyboardButton(text="🚀 Открыть тестовую версию", url=link)]]
                )
                await event.bot.send_message(uid, bot_link_text(link), reply_markup=kb)
        else:
            await event.bot.decline_chat_join_request(chat_id=event.chat.id, user_id=uid)
            logger.info("declined join request uid=%s chat=%s (not whitelisted)", uid, event.chat.id)
    except Exception:
        logger.exception("failed to handle join request uid=%s chat=%s", uid, event.chat.id)


if __name__ == "__main__":
    asyncio.run(main())
